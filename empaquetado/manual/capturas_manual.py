"""Pantallazos para el manual: recorre el programa (conectado a Render) con una cuenta y un contrato FICTICIOS.

Uso: python capturas_manual.py <escritorio> <backend> <carpeta_salida>
Marca con recuadros numerados los botones que el manual pide pulsar. Borra la cuenta de ejemplo al final.
"""
import hashlib
import os
import random
import secrets
import sys
import time
import traceback
from pathlib import Path

os.environ["QT_QPA_PLATFORM"] = "windows"
os.environ["PORTAL_KEYRING_SERVICIO"] = "PortalTrabajadorPruebas"  # nunca la sesión real del usuario
ESC, BACK, OUT = (Path(a).resolve() for a in sys.argv[1:4])
OUT.mkdir(parents=True, exist_ok=True)
os.environ["PORTAL_CARPETA_EVIDENCIAS"] = str(OUT / "evidencias_locales")
sys.path.insert(0, str(ESC))
os.chdir(ESC)

import httpx  # noqa: E402
from docx import Document  # noqa: E402
from dotenv import dotenv_values  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

env = dotenv_values(BACK / ".env")
URL, SECRETA = env["SUPABASE_URL"], env["SUPABASE_SERVICE_ROLE_KEY"]
ADMIN = {"apikey": SECRETA, "Authorization": f"Bearer {SECRETA}"}
CLAVE = "Clave-Ejemplo-2026"
NARANJA = (232, 89, 12)


def crear_usuario(email):
    return httpx.post(f"{URL}/auth/v1/admin/users", headers=ADMIN, timeout=30, json={
        "email": email, "password": CLAVE, "email_confirm": True,
        "user_metadata": {"nombre_completo": "Contratista de Ejemplo"}})


email = "usuario.ejemplo@example.com"
r = crear_usuario(email)
if r.status_code >= 400:
    email = f"usuario.ejemplo.{secrets.token_hex(2)}@example.com"
    r = crear_usuario(email)
uid = r.json()["id"]
cedula = str(random.randint(1_010_000_000, 1_019_999_999))
print("cuenta de ejemplo:", email)

MINUTA = (
    "ANEXO TÉCNICO No. 4521 DEL CONTRATO ELECTRONICO SECOP II. Entre los suscritos: FUNCIONARIO DE EJEMPLO, "
    "actuando en nombre y representación del Distrito Turístico y Cultural de Cartagena de Indias, con NIT "
    "890480184-4, por una parte; y por la otra, CONTRATISTA DE EJEMPLO, identificado(a) con la Cédula de Ciudadanía "
    f"No. {int(cedula):,}".replace(",", ".") + ", actuando en su propio nombre. I. Que existe insuficiencia de "
    "personal de planta para cumplir las funciones propias de la OFICINA DE EJEMPLO. CLÁUSULA 1 - OBJETO DEL "
    "CONTRATO: PRESTACIÓN DE SERVICIOS DE APOYO A LA GESTIÓN TECNOLÓGICA EN LA OFICINA DE EJEMPLO. CLÁUSULA 2 - "
    "OBLIGACIONES ESPECÍFICAS DEL CONTRATO: 1. REALIZAR SOPORTE TÉCNICO A LOS EQUIPOS TECNOLÓGICOS, DE "
    "TELECOMUNICACIONES Y SISTEMAS DE INFORMACIÓN DE LA DEPENDENCIA; 2. REALIZAR EL MANTENIMIENTO PREVENTIVO Y "
    "CORRECTIVO DE LOS EQUIPOS DE CÓMPUTO E IMPRESORAS; 3. ACTUALIZAR EL INVENTARIO DE LOS EQUIPOS TECNOLÓGICOS DE "
    "LA DEPENDENCIA; 4. CAPACITAR A LOS FUNCIONARIOS EN EL USO DE LAS HERRAMIENTAS TECNOLÓGICAS; 5. ASISTIR A LAS "
    "REUNIONES A LAS QUE SE LE CONVOQUE; 6. LAS DEMÁS QUE LE SEAN ASIGNADAS DE ACUERDO CON LA NATURALEZA DEL "
    "CONTRATO.. PARÁGRAFO PRIMERO: Serán de propiedad de la Entidad los resultados. CLÁUSULA 4 - VALOR DEL CONTRATO "
    "Y FORMA DE PAGO: El valor del presente Contrato corresponde a la suma de: NUEVE MILLONES DE PESOS MCTE "
    "($9,000,000.00). CLÁUSULA 6 - PLAZO DE EJECUCIÓN: El plazo de ejecución del presente Contrato será de TRES (3) "
    "meses. PARÁGRAFO: nada. CLÁUSULA 8 - OBLIGACIONES GENERALES DEL CONTRATISTA: El Contratista se obliga a: "
    "1. Cumplir con lo pactado. 2. Asistir a las reuniones. 3. Estar afiliado al sistema de seguridad social. "
    "4. Responder por los insumos. PARÁGRAFO: En los casos... CLÁUSULA 10 - OBLIGACIONES GENERALES DEL CONTRATANTE: "
    "1. Ejercer el control. 2. Pagar. CLÁUSULA 22 - SUPERVISIÓN: La supervisión de la ejecución estará a cargo del "
    "JEFE DE LA OFICINA DE EJEMPLO. CLÁUSULA 23 - ANEXOS.")
ruta_minuta = OUT / "Minuta de ejemplo.docx"
doc = Document()
doc.add_paragraph(MINUTA)
doc.save(ruta_minuta)

ruta_foto = OUT / "mantenimiento_equipos.jpg"
img = Image.new("RGB", (2400, 1600), (196, 205, 214))
dib = ImageDraw.Draw(img)
dib.rectangle([0, 1100, 2400, 1600], fill=(120, 92, 66))
for i in range(4):
    x = 180 + i * 560
    dib.rectangle([x, 520, x + 420, 860], fill=(28, 30, 36), outline=(90, 90, 95), width=10)
    dib.rectangle([x + 30, 550, x + 390, 830], fill=(52, 120, 168))
    dib.rectangle([x + 180, 860, x + 240, 960], fill=(60, 60, 64))
    dib.rectangle([x + 90, 960, x + 330, 990], fill=(60, 60, 64))
    dib.rectangle([x + 40, 1010, x + 380, 1070], fill=(225, 225, 228))
img.save(ruta_foto, "JPEG", quality=90)

from PySide6.QtCore import QPoint, Qt  # noqa: E402
from PySide6.QtWidgets import QApplication, QPushButton, QScrollArea  # noqa: E402

app = QApplication([])
from main import cargar_fuentes  # noqa: E402

cargar_fuentes()
from app import config, tema  # noqa: E402
from app.estado import estado  # noqa: E402

tema.aplicar(app, "claro")
from app.ventanas.principal import VentanaPrincipal  # noqa: E402

print("servidor:", config.API_BASE_URL, "| modo:", config.MODO_BACKEND)


def esperar_ms(ms):
    fin = time.time() + ms / 1000
    while time.time() < fin:
        QApplication.processEvents()
        time.sleep(0.01)


def esperar(cond, timeout=60, nombre=""):
    fin = time.time() + timeout
    while time.time() < fin:
        esperar_ms(100)
        try:
            if cond():
                esperar_ms(400)
                return True
        except Exception:
            pass
    raise TimeoutError(f"tiempo agotado esperando {nombre}")


v = VentanaPrincipal()
v.setAttribute(Qt.WA_DontShowOnScreen)
v.resize(1280, 800)
v.show()


def boton(texto, por="texto"):
    for b in v.findChildren(QPushButton):
        valor = b.text().strip() if por == "texto" else b.accessibleName()
        if valor == texto and b.isVisibleTo(v):
            return b
    raise LookupError(texto)


def visible(w):
    p = w.parentWidget()
    while p is not None:
        if isinstance(p, QScrollArea):
            p.ensureWidgetVisible(w, 40, 40)
        p = p.parentWidget()
    esperar_ms(150)


def foto(nombre, marcas=()):
    for w in marcas:
        visible(w)
    pix = v.grab()
    dpr = pix.devicePixelRatio()
    ruta = OUT / f"{nombre}.png"
    pix.save(str(ruta))
    if marcas:
        imagen = Image.open(ruta).convert("RGB")
        d = ImageDraw.Draw(imagen)
        letra = ImageFont.truetype("arialbd.ttf", int(23 * dpr))
        for n, w in enumerate(marcas, 1):
            p = w.mapTo(v, QPoint(0, 0))
            x0, y0 = (p.x() - 6) * dpr, (p.y() - 6) * dpr
            x1, y1 = (p.x() + w.width() + 6) * dpr, (p.y() + w.height() + 6) * dpr
            d.rounded_rectangle([x0, y0, x1, y1], radius=int(9 * dpr), outline=NARANJA, width=int(4 * dpr))
            radio = 18 * dpr
            cx, cy = max(x0, radio + 2), max(y0, radio + 2)
            d.ellipse([cx - radio, cy - radio, cx + radio, cy + radio], fill=NARANJA, outline="white", width=2)
            d.text((cx, cy), str(n), fill="white", font=letra, anchor="mm")
        imagen.save(ruta)
    print("  foto", nombre)


try:
    esperar(lambda: v.pila.currentWidget() is v.pag_login, 180, "login")

    # 00 · pantalla de carga mientras el servidor enciende
    v.pila.setCurrentWidget(v.pag_cargando)
    v.pag_cargando.empezar("Conectando con el servidor…")
    v.pag_cargando._inicio -= 18
    v.pag_cargando._tic()
    esperar_ms(300)
    foto("00_cargando")
    v.mostrar_login()
    esperar_ms(300)

    # 01 · crear cuenta (la cuenta real se creó por administración: no se envían correos)
    lg = v.pag_login
    lg.b_crear.click()
    esperar_ms(300)
    lg.nombre.poner("Contratista de Ejemplo")
    lg.correo.poner(email)
    lg.clave.poner(CLAVE)
    lg.acepto.setChecked(True)
    foto("01_crear_cuenta", [lg.b_crear, lg.nombre, lg.correo, lg.clave, lg.acepto, lg.enviar])
    lg.acepto.setChecked(False)

    # 02 · iniciar sesión
    lg.b_entrar.click()
    esperar_ms(300)
    lg.correo.poner(email)
    lg.clave.poner(CLAVE)
    foto("02_iniciar_sesion", [lg.correo, lg.clave, lg.enviar, lg.olvido])
    lg._enviar()
    esperar(lambda: v.pila.currentWidget() is v.pag_minuta, 180, "minuta")

    # 03 · subir la minuta
    mi = v.pag_minuta
    mi.zona._elegir(str(ruta_minuta))
    esperar_ms(300)
    foto("03_minuta", [mi.zona, mi.leer])
    mi._leer()
    esperar(lambda: v.pila.currentWidget() is v.pag_revisar, 240, "revisar")

    # 04 · revisar y confirmar
    rv = v.pag_revisar
    rv.fecha_contrato.poner("2026-09-25")
    rv.fecha_inicio.poner("2026-10-01")
    rv.fecha_fin.poner("2026-12-31")
    rv.sup_nombre.poner("NOMBRE DEL SUPERVISOR")
    rv.sup_cargo.poner("JEFE DE LA OFICINA DE EJEMPLO")
    foto("04_revisar_datos", [rv.nombre, rv.cedula])
    foto("04b_revisar_confirmar", [rv.guardar])
    rv._guardar()
    ini = v.secciones["inicio"]
    esperar(lambda: v.pila.currentWidget() is v.pag_app and not ini._viejo, 120, "inicio")

    # 05 · inicio y menú
    foto("05_inicio", [boton("Registrar actividad"), boton("Mi informe"), boton("Generar informe"),
                       boton("Planes"), v.medidor])

    # 06 · registrar actividad con evidencia
    act = v.secciones["actividad"]
    v.ir("actividad")
    esperar_ms(300)
    act.texto.poner("hoy hice mantenimiento a los computadores de la oficina")
    act.agregar_evidencias([str(ruta_foto)])
    esperar_ms(600)
    foto("06_actividad", [act.texto, act.fecha, boton("Adjuntar evidencia"), act.b_ubicar])
    act._clasificar()
    esperar(lambda: act.pila.currentWidget() is act.prop_pagina, 120, "propuesta")

    # 07 · propuesta de la IA
    pp = act.propuesta
    foto("07_propuesta", [pp.combo, pp.redaccion, pp.b_guardar])
    act._guardar()
    esperar(lambda: act.pila.currentWidget() is act.form and act.aviso.isVisible(), 180, "guardado")
    foto("08_actividad_guardada", [act.aviso])

    # segunda actividad, sin pantallazo, para que el informe se vea completo
    act.texto.poner("capacité a los funcionarios en el uso del correo institucional")
    act._clasificar()
    esperar(lambda: act.pila.currentWidget() is act.prop_pagina, 120, "propuesta 2")
    act._guardar()
    esperar(lambda: act.pila.currentWidget() is act.form and act.aviso.isVisible(), 180, "guardado 2")

    # 09 · mi informe
    v.ir("informe")
    inf = v.secciones["informe"]
    esperar(lambda: not inf._viejo and inf.datos, 120, "informe")
    foto("09_mi_informe")

    # 10 y 11 · generar y guardar
    v.ir("generar")
    gen = v.secciones["generar"]
    esperar(lambda: not gen._viejo and gen.datos, 120, "generar")
    foto("10_generar", [gen.selector, gen.final, gen.fecha, gen.b_generar])
    gen._generar()
    esperar(lambda: gen.visor.isVisible(), 300, "vista previa")
    esperar_ms(1500)
    foto("11_generado", [gen.b_word, gen.b_pdf])
    _, pdf = estado().api.descargar_informe(gen.datos["informe"]["id"], "pdf")
    import pypdfium2 as pdfium
    pdoc = pdfium.PdfDocument(pdf)
    pdoc[0].render(scale=1.4).to_pil().save(OUT / "20_informe_pagina1.png")
    pdoc[len(pdoc) - 1].render(scale=1.4).to_pil().save(OUT / "21_informe_fotos.png")
    print("  páginas del informe:", len(pdoc))

    # 12 y 13 · planes y espera del pago (sin pagar: nunca se abre Wompi)
    v.ir("planes")
    pl = v.secciones["planes"]
    esperar(lambda: bool(pl.planes), 120, "planes")
    esperar_ms(500)
    foto("12_planes", [boton("Pagar plan Normal", por="nombre")])
    pl.espera.show()
    esperar_ms(300)
    foto("13_esperando_pago", [pl.verificar])
    pl.espera.hide()

    # 14 · perfil
    v.ir("perfil")
    perf = v.secciones["perfil"]
    esperar_ms(3000)
    foto("14_perfil", [perf.cedula, boton("Cargar un contrato nuevo")])
except Exception:
    traceback.print_exc()
finally:
    try:
        estado().gestor.establecer(None)
    except Exception:
        pass
    tok = httpx.post(f"{URL}/auth/v1/token?grant_type=password", headers={"apikey": env["SUPABASE_ANON_KEY"]},
                     json={"email": email, "password": CLAVE}, timeout=30).json().get("access_token")
    borrada = False
    if tok:
        r = httpx.delete(f"{config.API_BASE_URL}/api/cuenta", headers={"Authorization": f"Bearer {tok}"},
                         timeout=180)
        borrada = r.status_code == 200
    if not borrada:
        httpx.delete(f"{URL}/auth/v1/admin/users/{uid}", headers=ADMIN, timeout=30)
    httpx.delete(f"{URL}/rest/v1/pruebas_gratis", headers=ADMIN, timeout=30,
                 params={"cedula_hash": f"eq.{hashlib.sha256(cedula.encode()).hexdigest()}"})
    print("cuenta de ejemplo eliminada")

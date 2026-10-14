"""Manual de uso con pantallazos (Word -> PDF con el Word instalado).

Uso: python crear_manual.py <backend> <carpeta_capturas> <icono.png> <pdf_salida>
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

BACKEND, CAP, ICONO, SALIDA = (Path(a) for a in sys.argv[1:5])
sys.path.insert(0, str(BACKEND.resolve()))
from servicios.conversor_pdf import convertir  # noqa: E402

AZUL = RGBColor(0x0E, 0x5A, 0x6B)
NARANJA = RGBColor(0xE8, 0x59, 0x0C)
GRIS = RGBColor(0x55, 0x5F, 0x6B)
ANCHO = Inches(6.4)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.95)
sec.top_margin = Inches(0.8)
sec.bottom_margin = Inches(0.75)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(5)
for nombre, tam in (("Title", 30), ("Heading 1", 17), ("Heading 2", 13)):
    s = doc.styles[nombre]
    s.font.name = "Calibri"
    s.font.size = Pt(tam)
    s.font.color.rgb = AZUL
    s.font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(6)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)

# Pie de página: nombre del manual y número de página
pie = sec.footer.paragraphs[0]
pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = pie.add_run("Portal del Trabajador · Manual de uso · página ")
r.font.size = Pt(9)
r.font.color.rgb = GRIS
campo = OxmlElement("w:fldSimple")
campo.set(qn("w:instr"), "PAGE")
rr = OxmlElement("w:r")
rpr = OxmlElement("w:rPr")
tam = OxmlElement("w:sz")
tam.set(qn("w:val"), "18")
rpr.append(tam)
rr.append(rpr)
t = OxmlElement("w:t")
t.text = "1"
rr.append(t)
campo.append(rr)
pie._p.append(campo)


def p(texto="", negrita=None, despues=None, color=None, tam=None, italica=False, centro=False):
    par = doc.add_paragraph()
    if negrita:
        par.add_run(negrita).bold = True
    run = par.add_run(texto)
    run.italic = italica
    if color:
        run.font.color.rgb = color
    if tam:
        run.font.size = Pt(tam)
    if despues is not None:
        par.paragraph_format.space_after = Pt(despues)
    if centro:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return par


def vinetas(items, numerada=False):
    for item in items:
        par = doc.add_paragraph(style="List Number" if numerada else "List Bullet")
        if isinstance(item, tuple):
            par.add_run(item[0]).bold = True
            par.add_run(item[1])
        else:
            par.add_run(item)
        par.paragraph_format.space_after = Pt(3)


def pantallazo(archivo, pie_figura, pasos=(), recorte=None, ancho=ANCHO):
    """Imagen con borde fino, su pie y la explicación de cada número marcado."""
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.page_break_before = consumir_salto()
    ruta = CAP / archivo
    if recorte:  # fracciones (izq, arriba, der, abajo) de la imagen
        img = Image.open(ruta)
        w, h = img.size
        ruta = CAP / f"_recorte_{archivo}"
        img.crop((int(recorte[0] * w), int(recorte[1] * h), int(recorte[2] * w), int(recorte[3] * h))).save(ruta)
    run = par.add_run()
    run.add_picture(str(ruta), width=ancho)
    borde = OxmlElement("a:ln")
    borde.set("w", "9525")
    relleno = OxmlElement("a:solidFill")
    color = OxmlElement("a:srgbClr")
    color.set("val", "C5CCD3")
    relleno.append(color)
    borde.append(relleno)
    sppr = run._r.find(".//" + qn("pic:spPr"))
    if sppr is not None:
        sppr.append(borde)
    leyenda = p(pie_figura, italica=True, color=GRIS, tam=9, despues=6 if pasos else 12, centro=True)
    leyenda.paragraph_format.keep_with_next = bool(pasos)
    for i, (titulo, texto) in enumerate(pasos, 1):
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.32)
        par.paragraph_format.first_line_indent = Inches(-0.32)
        par.paragraph_format.space_after = Pt(3 if i < len(pasos) else 12)
        par.paragraph_format.keep_with_next = i < len(pasos)
        num = par.add_run(f"{i}\t")
        num.bold = True
        num.font.color.rgb = NARANJA
        par.paragraph_format.tab_stops.add_tab_stop(Inches(0.32))
        par.add_run(titulo).bold = True
        par.add_run(texto)


def recuadro(texto, titulo="Importante"):
    tabla = doc.add_table(rows=1, cols=1)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = tabla.rows[0].cells[0]
    tcpr = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), "E7F1F3")
    tcpr.append(sombra)
    bordes = OxmlElement("w:tcBorders")
    for lado in ("top", "bottom", "right"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "nil")
        bordes.append(b)
    izq = OxmlElement("w:left")
    izq.set(qn("w:val"), "single")
    izq.set(qn("w:sz"), "24")
    izq.set(qn("w:color"), "0E5A6B")
    bordes.append(izq)
    tcpr.append(bordes)
    par = celda.paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run(f"{titulo}: ")
    r.bold = True
    r.font.color.rgb = AZUL
    par.add_run(texto)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


_SALTO = [False]


def salto():
    """El siguiente título o imagen empieza en página nueva (sin párrafos vacíos que dejen hojas en blanco)."""
    _SALTO[0] = True


def consumir_salto() -> bool:
    valor, _SALTO[0] = _SALTO[0], False
    return valor


def titulo_seccion(texto, level=1):
    par = doc.add_heading(texto, level=level)
    par.paragraph_format.page_break_before = consumir_salto()
    return par


# ---------------------------------------------------------------- Portada
for _ in range(3):
    doc.add_paragraph()
portada = doc.add_paragraph()
portada.alignment = WD_ALIGN_PARAGRAPH.CENTER
portada.add_run().add_picture(str(ICONO), width=Inches(1.1))
titulo = doc.add_paragraph("Portal del Trabajador", style="Title")
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Manual de uso", tam=18, color=AZUL, centro=True, despues=2)
p("Informe mensual de actividades · Formato GDOGC01-F007 (Informe de Supervisión)", tam=11, color=GRIS,
  centro=True, despues=26)
p("Escribe con tus palabras lo que hiciste cada día. El programa ubica cada actividad en la obligación de tu "
  "contrato, la redacta en lenguaje formal y, al final del mes, te entrega el informe listo en Word y PDF, con "
  "tus fotos de evidencia como anexos.", centro=True, despues=30)
p("Versión 1.1.0 · octubre de 2026", tam=10, color=GRIS, centro=True)
salto()

titulo_seccion("Contenido", level=1)
for linea in ("1. Antes de empezar", "2. Instalar el programa", "3. Abrir el programa",
              "4. Crear tu cuenta", "5. Iniciar sesión", "6. Cargar tu contrato (una sola vez)",
              "7. La pantalla de inicio", "8. Registrar una actividad", "9. Revisar tu informe del mes",
              "10. Generar y descargar el informe", "11. Planes y pagos", "12. Tu perfil",
              "13. Preguntas frecuentes"):
    p(linea, despues=3)
p("")
recuadro("en los pantallazos, los recuadros naranjas con un número señalan dónde hacer clic o qué revisar. "
         "Debajo de cada imagen se explica cada número.", "Cómo leer este manual")
recuadro("los datos que aparecen en las imágenes (nombre, cédula, contrato) son de ejemplo.", "Nota")
salto()

# ---------------------------------------------------------------- 1 y 2
titulo_seccion("1. Antes de empezar", level=1)
vinetas(["Un computador con Windows 10 u 11 y conexión a internet.",
         "Tu minuta (el documento de tu contrato) en PDF o Word.",
         "Un correo electrónico al que tengas acceso.",
         "No necesitas permisos de administrador para instalar el programa."])

titulo_seccion("2. Instalar el programa", level=1)
vinetas([("Abre el instalador: ", "haz doble clic en «PortalTrabajador-Instalador-1.1.0.exe»."),
         ("Si Windows muestra «Windows protegió su PC»: ", "haz clic en «Más información» y luego en «Ejecutar de "
          "todas formas». Este aviso aparece porque el programa es nuevo y todavía no tiene firma digital; "
          "es seguro continuar."),
         ("Sigue el asistente: ", "pulsa «Siguiente» y luego «Instalar». Si quieres, marca la opción de crear un "
          "acceso directo en el escritorio."),
         ("Abre el programa: ", "al terminar, deja marcada la casilla «Abrir Portal del Trabajador» y pulsa "
          "«Finalizar». Después lo encuentras en el escritorio o en el menú Inicio.")], numerada=True)
recuadro("el programa se instala solo para tu usuario de Windows. Para desinstalarlo, ve a Configuración → "
         "Aplicaciones y busca «Portal del Trabajador».", "Dato")

# ---------------------------------------------------------------- 3
titulo_seccion("3. Abrir el programa", level=1)
p("Al abrirlo, el programa se conecta al servidor. Normalmente tarda unos segundos. Si nadie lo ha usado en un "
  "rato, el servidor estaba en reposo y puede tardar hasta un minuto en encender: verás este mensaje con un "
  "contador de segundos. No cierres el programa; entra solo cuando el servidor responde.")
pantallazo("00_cargando.png", "Figura 1. Pantalla de carga mientras el servidor enciende.",
           recorte=(0.25, 0.36, 0.75, 0.64), ancho=Inches(4.6))
salto()

# ---------------------------------------------------------------- 4
titulo_seccion("4. Crear tu cuenta", level=1)
p("Solo la primera vez. Cada persona tiene una sola cuenta, vinculada a su número de cédula.")
pantallazo("01_crear_cuenta.png", "Figura 2. Crear una cuenta nueva.", [
    ("Crear cuenta: ", "pulsa esta pestaña."),
    ("Nombre completo: ", "escríbelo como aparece en tu contrato."),
    ("Correo: ", "uno al que tengas acceso; ahí llegan los avisos de tu cuenta."),
    ("Contraseña: ", "de 8 caracteres o más. Guárdala bien."),
    ("Acepto el tratamiento de mis datos: ", "marca la casilla (puedes leer el aviso en «Leer aviso»)."),
    ("Crear cuenta: ", "pulsa el botón para terminar.")])
recuadro("revisa tu correo (también la carpeta de spam) y abre el enlace de confirmación. Si después el "
         "navegador muestra una página de error, no te preocupes: tu cuenta ya quedó confirmada. Vuelve al "
         "programa e inicia sesión.", "Confirma tu correo")
salto()

# ---------------------------------------------------------------- 5
titulo_seccion("5. Iniciar sesión", level=1)
pantallazo("02_iniciar_sesion.png", "Figura 3. Iniciar sesión.", [
    ("Correo: ", "el de tu cuenta."),
    ("Contraseña: ", "la que creaste."),
    ("Entrar: ", "pulsa para ingresar. El programa recuerda tu sesión en este computador."),
    ("¿Olvidaste tu contraseña?: ", "te enviamos un código a tu correo para crear una nueva.")])
salto()

# ---------------------------------------------------------------- 6
titulo_seccion("6. Cargar tu contrato (una sola vez)", level=1)
p("La primera vez que entras, el programa te pide la minuta de tu contrato. De ahí saca tus obligaciones, las "
  "fechas y los datos que van en el informe.")
titulo_seccion("Paso 1 · Sube tu minuta", level=2)
pantallazo("03_minuta.png", "Figura 4. Subir la minuta.", [
    ("Zona del archivo: ", "haz clic para buscar tu minuta (PDF o Word) o arrástrala hasta aquí."),
    ("Leer mi contrato: ", "pulsa y espera unos segundos mientras se leen tus datos y obligaciones.")])
salto()
titulo_seccion("Paso 2 · Revisa tus datos", level=2)
p("Revisa con calma lo que se leyó y corrige lo que haga falta. Los avisos amarillos de arriba te dicen qué datos "
  "no venían en la minuta (por ejemplo, las fechas o el nombre del supervisor): escríbelos.")
pantallazo("04_revisar_datos.png", "Figura 5. Revisión de los datos leídos de la minuta.", [
    ("Nombre completo: ", "como aparece en el contrato."),
    ("Número de cédula: ", "debe ser exactamente la que aparece en tu minuta. Queda vinculada a tu cuenta y no se "
     "puede cambiar después.")])
p("Más abajo encontrarás los datos del contrato (número, objeto, dependencia, fechas, valor), los del supervisor y "
  "la lista de obligaciones específicas. Puedes corregir el texto de cualquier obligación, agregar una que falte "
  "o borrar una que sobre con el ícono de la papelera.")
salto()
titulo_seccion("Paso 3 · Confirma", level=2)
pantallazo("04b_revisar_confirmar.png", "Figura 6. Final de la revisión.", [
    ("Confirmar y entrar: ", "guarda tu contrato y te lleva a la pantalla de inicio. Si te equivocaste de "
     "archivo, usa «Volver a subir la minuta».")])
recuadro("si la cédula ya está registrada en otra cuenta, el programa no te dejará guardar el contrato. Cada "
         "persona solo puede tener una cuenta.", "Una cuenta por persona")
salto()

# ---------------------------------------------------------------- 7
titulo_seccion("7. La pantalla de inicio", level=1)
p("Aquí ves el resumen de tu mes: actividades registradas, obligaciones cubiertas, informes disponibles y los días "
  "que faltan para cerrar el periodo. También puedes escribir una actividad rápida en «¿Qué hiciste hoy?».")
pantallazo("05_inicio.png", "Figura 7. Pantalla de inicio y menú.", [
    ("Registrar actividad: ", "para anotar lo que hiciste, con evidencias."),
    ("Mi informe: ", "para ver y corregir las actividades del mes, ordenadas por obligación."),
    ("Generar informe: ", "para crear el informe en Word y PDF."),
    ("Planes: ", "para ver los planes y pagar."),
    ("Tu cupo: ", "cuántos informes te quedan en tu plan. Haz clic para ir a Planes.")])
salto()

# ---------------------------------------------------------------- 8
titulo_seccion("8. Registrar una actividad", level=1)
p("Hazlo el mismo día, en una frase sencilla. No tienes que escribir formal: el programa se encarga de eso.")
pantallazo("06_actividad.png", "Figura 8. Registrar una actividad con evidencia.", [
    ("¿Qué hiciste?: ", "escríbelo con tus palabras. Ejemplo: «hoy hice mantenimiento a los computadores de la "
     "oficina». Una actividad por registro."),
    ("Fecha de la actividad: ", "viene con la fecha de hoy; cámbiala si la actividad fue otro día."),
    ("Adjuntar evidencia (opcional): ", "fotos (JPG o PNG), actas, listados en PDF, Word o Excel. Se guarda una "
     "copia en Documentos › Portal del Trabajador › Evidencias."),
    ("Ubicar en mi informe: ", "el programa busca a qué obligación corresponde y la redacta en lenguaje formal. "
     "Si prefieres hacerlo tú, usa «Guardar sin analizar».")])
salto()
titulo_seccion("Revisa la propuesta y guarda", level=2)
pantallazo("07_propuesta.png", "Figura 9. Propuesta: obligación y redacción formal.", [
    ("Obligación: ", "la obligación de tu contrato en la que encaja la actividad. Si no es la correcta, elige "
     "otra en la lista."),
    ("Redacción sugerida: ", "el texto formal que irá en el informe. Puedes editarlo libremente. El programa no "
     "inventa datos: solo mejora la redacción de lo que escribiste."),
    ("Guardar: ", "guarda la actividad. «Descartar» la borra sin guardar.")])
p("Si la confianza es media o baja, el programa te muestra opciones y puede hacerte una pregunta para ubicarla "
  "mejor. Siempre decides tú.", italica=True, color=GRIS)
salto()
pantallazo("08_actividad_guardada.png", "Figura 10. Actividad guardada.", [
    ("Confirmación: ", "indica en qué obligación quedó la actividad y la descripción del anexo de la evidencia "
     "(«Registro fotográfico de…»). En el informe la actividad termina con «Se adjunta como evidencia el "
     "Anexo N» y la foto se pega al final del documento.")])
salto()

# ---------------------------------------------------------------- 9
titulo_seccion("9. Revisar tu informe del mes", level=1)
p("En «Mi informe» ves todas las actividades del periodo, agrupadas por obligación, tal como saldrán en el "
  "informe. Cambia el mes en «Periodo».")
pantallazo("09_mi_informe.png", "Figura 11. Mi informe: actividades por obligación.")
vinetas([("Lápiz: ", "editar la redacción, la fecha o las evidencias de la actividad (también la descripción de "
          "cada anexo)."),
         ("Flechas: ", "mover la actividad a otra obligación."),
         ("Papelera: ", "eliminar la actividad."),
         ("Avisos amarillos: ", "te recuerdan qué obligaciones aún no tienen actividades en el mes.")])
salto()

# ---------------------------------------------------------------- 10
titulo_seccion("10. Generar y descargar el informe", level=1)
pantallazo("10_generar.png", "Figura 12. Antes de generar.", [
    ("Periodo: ", "el mes del informe."),
    ("Marcar como informe final: ", "márcalo solo en el último informe de tu contrato. Si no, sale como parcial."),
    ("Fecha de elaboración: ", "la fecha que aparecerá en el informe."),
    ("Generar informe: ", "crea el informe con el formato oficial. Tarda unos segundos.")])
salto()
pantallazo("11_generado.png", "Figura 13. Vista previa del informe generado.", [
    ("Guardar Word: ", "guarda el informe editable (.docx)."),
    ("Guardar PDF: ", "guarda el informe listo para firmar y enviar.")])
recuadro("registrar actividades no gasta cupo; solo generar el informe. Si corriges algo y vuelves a generar el "
         "informe del mismo mes, no se descuenta otro.", "Cupo")
salto()
titulo_seccion("Así queda tu informe", level=2)
tabla = doc.add_table(rows=2, cols=2)
tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (archivo, texto) in enumerate((("20_informe_pagina1.png", "Primera página: datos del contrato y tabla de "
                                       "actividades por obligación."),
                                      ("21_informe_fotos.png", "Última página: registro fotográfico de los anexos."))):
    celda = tabla.rows[0].cells[i]
    par = celda.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(CAP / archivo), width=Inches(3.05))
    leyenda = tabla.rows[1].cells[i].paragraphs[0]
    leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = leyenda.add_run(texto)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GRIS
salto()

# ---------------------------------------------------------------- 11
titulo_seccion("11. Planes y pagos", level=1)
planes = doc.add_table(rows=1, cols=3)
planes.style = "Table Grid"
planes.alignment = WD_TABLE_ALIGNMENT.CENTER
for celda, texto in zip(planes.rows[0].cells, ("Plan", "Precio", "Qué incluye")):
    celda.text = ""
    r = celda.paragraphs[0].add_run(texto)
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), "0E5A6B")
    celda._tc.get_or_add_tcPr().append(sombra)
for fila in (("Gratis (prueba)", "$0", "1 informe en total. Una sola vez por cédula."),
             ("Normal", "$15.000 por 30 días", "2 informes en el periodo de 30 días."),
             ("Pro", "$50.000 por 30 días", "Informes ilimitados durante 30 días.")):
    for celda, texto in zip(planes.add_row().cells, fila):
        celda.text = texto
p("")
pantallazo("12_planes.png", "Figura 14. Planes.", [
    ("Pagar: ", "se abre en tu navegador la página segura de Wompi (Bancolombia). Paga con tarjeta, Nequi, PSE u "
     "otro medio de los que te muestre. El programa nunca ve los datos de tu tarjeta.")])
salto()
pantallazo("13_esperando_pago.png", "Figura 15. Esperando la confirmación del pago.", [
    ("Ya pagué, verificar: ", "cuando termines de pagar, vuelve al programa. El plan se activa solo en unos "
     "segundos; si no, pulsa este botón para consultar el pago.")])
recuadro("el plan dura 30 días desde el pago. Antes de que venza, puedes renovarlo desde esta misma pantalla.",
         "Vigencia")
salto()

# ---------------------------------------------------------------- 12
titulo_seccion("12. Tu perfil", level=1)
pantallazo("14_perfil.png", "Figura 16. Perfil y suscripción.", [
    ("Número de cédula: ", "vinculada a tu cuenta desde tu contrato; no se puede cambiar."),
    ("Cargar un contrato nuevo: ", "úsalo cuando firmes un contrato nuevo. Tus informes anteriores no se "
     "borran.")])
vinetas([("Editar datos del contrato: ", "corrige fechas, supervisor u obligaciones."),
         ("Cerrar sesión: ", "en el menú de la izquierda, abajo."),
         ("Eliminar mi cuenta: ", "al final del perfil. Borra tus datos y archivos (Ley 1581 de 2012).")])
salto()

# ---------------------------------------------------------------- 13
titulo_seccion("13. Preguntas frecuentes", level=1)
preguntas = [
    ("El programa dice «Conectando con el servidor…» por un rato.",
     "Es normal si nadie lo había usado en un rato: el servidor tarda hasta un minuto en encender. Espera sin "
     "cerrar el programa."),
    ("Dice «No pudimos conectarnos con el servidor».",
     "Revisa tu conexión a internet y vuelve a abrir el programa."),
    ("La actividad no se pudo analizar.",
     "Intenta de nuevo en unos minutos, o usa «Guardar sin analizar» y elige tú la obligación."),
    ("Me dice que mi cédula ya está registrada.",
     "Ya existe una cuenta con esa cédula. Entra con esa cuenta; si no la reconoces, comunícate con soporte."),
    ("Pagué y mi plan no cambió.",
     "Vuelve a «Planes» y pulsa «Ya pagué, verificar». Si el pago fue aprobado, el plan se activa."),
    ("¿Dónde quedan mis fotos de evidencia?",
     "En Documentos › Portal del Trabajador › Evidencias, ordenadas por mes. También en la pantalla «Registrar "
     "actividad» con «Abrir carpeta de evidencias»."),
    ("Olvidé mi contraseña.",
     "En la pantalla de inicio de sesión pulsa «¿Olvidaste tu contraseña?» y sigue los pasos con el código que "
     "llega a tu correo."),
    ("Necesito ayuda.",
     "Comunícate con quien te compartió el programa y envíale el archivo de registro: escribe "
     "%LOCALAPPDATA%\\PortalTrabajador\\Logs en la barra de direcciones del Explorador de archivos y adjunta "
     "«portal.log»."),
]
for pregunta, respuesta in preguntas:
    par = p(pregunta, despues=1)
    par.runs[0].bold = True
    par.runs[0].font.color.rgb = AZUL
    par.paragraph_format.keep_with_next = True
    p(respuesta, despues=8)

tmp = SALIDA.parent / "_manual_temporal.docx"
doc.save(tmp)
pdf = convertir(tmp.read_bytes(), timeout=300)
tmp.unlink()
if not pdf:
    sys.exit("No se pudo convertir el manual a PDF")
SALIDA.write_bytes(pdf)
print("manual:", SALIDA.name, f"{len(pdf) / 1024 / 1024:.1f} MB")

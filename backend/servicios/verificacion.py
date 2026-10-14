"""Verificación automática del informe generado (sección 8.4).

- Las fuentes y tamaños usados en el informe deben existir ya en la plantilla.
- Márgenes y tamaño de página deben ser los mismos.
- Se renderiza la primera página de la plantilla en blanco y del informe y se
  compara el encabezado (escudo y títulos), que no debía cambiar.
Las diferencias se devuelven como advertencias y se registran en el log.
"""
from __future__ import annotations

import io
import logging

from docx import Document

from servicios.docx_util import fuentes_usadas

log = logging.getLogger("portal.verificacion")

UMBRAL_ENCABEZADO = 6.0  # diferencia media de gris (0-255) tolerada en la franja del encabezado


def comparar_docx(plantilla: bytes, informe: bytes) -> list[str]:
    advertencias = []
    doc_p, doc_i = Document(io.BytesIO(plantilla)), Document(io.BytesIO(informe))
    nuevas = fuentes_usadas(doc_i.element.body) - fuentes_usadas(doc_p.element.body)
    # La X de las casillas usa Arial 9 en negrita: Arial ya está en la plantilla
    nuevas = {n for n in nuevas if n != ("tamaño", "9 pt")}
    if nuevas:
        advertencias.append("El informe usa formatos que no están en la plantilla: "
                            + ", ".join(f"{t} {v}" for t, v in sorted(nuevas)))
    for sp, si in zip(doc_p.sections, doc_i.sections):
        if (sp.page_width, sp.page_height, sp.left_margin, sp.right_margin, sp.top_margin, sp.bottom_margin) != \
                (si.page_width, si.page_height, si.left_margin, si.right_margin, si.top_margin, si.bottom_margin):
            advertencias.append("El tamaño de página o los márgenes cambiaron respecto a la plantilla.")
    return advertencias


def _primera_pagina(pdf: bytes):
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(pdf)
    pagina = doc[0]
    ancho, alto = pagina.get_size()
    imagen = pagina.render(scale=0.6).to_pil().convert("L")
    return (ancho, alto), imagen


def comparar_pdf(pdf_plantilla: bytes, pdf_informe: bytes) -> list[str]:
    try:
        from PIL import ImageChops, ImageStat
        (tam_p, img_p), (tam_i, img_i) = _primera_pagina(pdf_plantilla), _primera_pagina(pdf_informe)
    except Exception as exc:  # sin pypdfium2/Pillow no se compara, pero no se bloquea la entrega
        log.info("No se pudo renderizar para comparar: %s", type(exc).__name__)
        return []
    advertencias = []
    if abs(tam_p[0] - tam_i[0]) > 1 or abs(tam_p[1] - tam_i[1]) > 1:
        advertencias.append("El tamaño de página del PDF no coincide con el de la plantilla.")
        return advertencias
    w, h = img_p.size
    # Franja del encabezado: escudo y títulos (se deja fuera "Página X de Y", que sí cambia)
    caja = (0, 0, int(w * 0.70), int(h * 0.16))
    diferencia = ImageStat.Stat(ImageChops.difference(img_p.crop(caja), img_i.crop(caja))).mean[0]
    if diferencia > UMBRAL_ENCABEZADO:
        advertencias.append(f"El encabezado del PDF se ve distinto al de la plantilla (diferencia {diferencia:.1f}).")
    return advertencias


def verificar(plantilla: bytes, informe: bytes, pdf_plantilla: bytes | None, pdf_informe: bytes | None) -> list[str]:
    advertencias = comparar_docx(plantilla, informe)
    if pdf_plantilla and pdf_informe:
        advertencias += comparar_pdf(pdf_plantilla, pdf_informe)
    for a in advertencias:
        log.warning("Verificación: %s", a)
    return advertencias

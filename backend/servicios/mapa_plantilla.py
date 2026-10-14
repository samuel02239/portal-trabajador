"""Mapa de campos de una plantilla Word (sección 8.2, paso 2).

Primero se ubican los campos por sus etiquetas ("NOMBRE Y APELLIDOS DEL
CONTRATISTA", "FECHA DE INICIO", la tabla con columnas OBLIGACIONES y
ACTIVIDADES, el párrafo del periodo con espacios ___, etc.). Si falta algo
obligatorio y la IA está configurada, se le pide que ubique lo que falta.
El usuario confirma el mapa una sola vez y queda guardado.
"""
from __future__ import annotations

import io
import logging
import re
from typing import Optional

from docx import Document
from pydantic import BaseModel

import prompts
from ia import IAError, ia_configurada, llamar_ia
from servicios.docx_util import (celdas, filas, normalizar, parrafos_cuerpo, q, tablas_cuerpo, texto_de)

log = logging.getLogger("portal.mapa")

CAMPOS: dict[str, tuple[str, list[str]]] = {
    "nombre_contratista": ("Nombre del contratista", [
        "nombre y apellidos del contratista", "nombres y apellidos del contratista", "nombre del contratista",
        "nombre del practicante", "nombre del estudiante", "nombre completo", "contratista"]),
    "numero_contrato": ("Número de contrato", [
        "n de contrato", "no de contrato", "numero de contrato", "nro de contrato", "numero del contrato",
        "n del contrato", "contrato no", "numero del acto administrativo", "acto administrativo"]),
    "fecha_contrato": ("Fecha del contrato", [
        "fecha del contrato", "fecha de suscripcion", "fecha del acto administrativo", "fecha de firma"]),
    "objeto": ("Objeto del contrato", ["objeto del contrato", "objeto contractual", "objeto"]),
    "cedula": ("Cédula", ["cedula", "rut", "documento de identidad", "identificacion", "c c"]),
    "fecha_inicio": ("Fecha de inicio", ["fecha de inicio", "fecha inicio", "fecha de iniciacion"]),
    "fecha_fin": ("Fecha de terminación", [
        "fecha de terminacion", "fecha de finalizacion", "fecha de vencimiento", "fecha final", "fecha fin"]),
    "dependencia": ("Dependencia", ["dependencia", "area donde presta", "secretaria"]),
    "supervisor": ("Nombre y cargo del supervisor", ["nombre y cargo del supervisor", "supervisor", "tutor"]),
}
OBLIGATORIOS = ["nombre_contratista", "cedula", "objeto", "ejecucion"]
_BLANCO = re.compile(r"_{2,}")


def _vacia(tc) -> bool:
    return len(texto_de(tc).replace("\xa0", " ").strip()) < 2


def _mejor_campo(etiqueta: str) -> tuple[str | None, int]:
    etiqueta = f" {etiqueta} "
    mejor, puntos = None, 0
    for campo, (_, sinonimos) in CAMPOS.items():
        for s in sinonimos:
            if f" {s} " in etiqueta and len(s) > puntos:
                mejor, puntos = campo, len(s)
    return mejor, puntos


def _detectar_ejecucion(tablas) -> dict | None:
    for ti, tbl in enumerate(tablas):
        for ri, tr in enumerate(filas(tbl)[:5]):
            textos = [normalizar(texto_de(tc)) for tc in celdas(tr)]
            col_obl = next((i for i, t in enumerate(textos) if "obligacion" in t), None)
            col_act = next((i for i, t in enumerate(textos) if "actividad" in t and i != col_obl), None)
            if col_obl is None or col_act is None:
                continue
            col_item = next((i for i, t in enumerate(textos)
                             if i not in (col_obl, col_act) and (t.startswith("item") or t in ("no", "n", "numero"))),
                            None)
            if col_item is None and 0 not in (col_obl, col_act):
                col_item = 0
            inicio = ri + 1
            todas = filas(tbl)
            if inicio >= len(todas):
                continue
            formato = None
            primera = celdas(todas[inicio])
            if col_obl < len(primera):
                t = texto_de(primera[col_obl]).strip()
                if t.startswith("(") or "ejemplo" in normalizar(texto_de(todas[inicio])):
                    formato = inicio
            return {"tabla": ti, "fila_encabezado": ri, "fila_inicio": inicio, "col_item": col_item,
                    "col_obligacion": col_obl, "col_actividades": col_act, "fila_formato": formato,
                    "eliminar_sobrantes": True}
    return None


def _contar_cajas_vacias(parrafos) -> int:
    n = 0
    for p in parrafos:
        for ac in p.iter(q("mc:AlternateContent")):
            caja = ac.find(f".//{q('wps:txbx')}")
            if caja is not None and not texto_de(caja.find(q("w:txbxContent"))).strip():
                n += 1
    return n


def detectar(contenido: bytes) -> dict:
    """Mapa de campos por reglas."""
    doc = Document(io.BytesIO(contenido))
    body = doc.element.body
    tablas = tablas_cuerpo(body)
    parrafos = parrafos_cuerpo(body)

    ejecucion = _detectar_ejecucion(tablas)
    campos: dict[str, dict] = {}
    puntajes: dict[str, int] = {}
    for ti, tbl in enumerate(tablas):
        if ejecucion and ti == ejecucion["tabla"]:
            continue
        for ri, tr in enumerate(filas(tbl)):
            tcs = celdas(tr)
            for ci in range(len(tcs) - 1):
                etiqueta = normalizar(texto_de(tcs[ci]))
                if not etiqueta or len(etiqueta) > 90 or not _vacia(tcs[ci + 1]):
                    continue
                campo, puntos = _mejor_campo(etiqueta)
                if campo and puntos > puntajes.get(campo, 0):
                    campos[campo] = {"tabla": ti, "fila": ri, "col": ci + 1, "etiqueta": texto_de(tcs[ci]).strip()}
                    puntajes[campo] = puntos

    parr: dict[str, dict] = {}
    for i, p in enumerate(parrafos):
        crudo, norm = texto_de(p), normalizar(texto_de(p))
        if "fecha_elaboracion" not in parr and "fecha de elaboracion" in norm and _BLANCO.search(crudo):
            parr["fecha_elaboracion"] = {"indice": i, "texto": crudo.strip()[:80]}
        elif "tipo_informe" not in parr and ("informe parcial" in norm or ("parcial" in norm and "final" in norm)):
            parr["tipo_informe"] = {"indice": i, "texto": crudo.strip()[:80]}
        elif "periodo" not in parr and "periodo" in norm and len(_BLANCO.findall(crudo)) >= 3:
            parr["periodo"] = {"indice": i, "texto": crudo.strip()[:80]}

    anexos = None
    for i, p in enumerate(parrafos):
        if normalizar(texto_de(p)).startswith("anexos"):
            items = []
            for j in range(i + 1, min(i + 40, len(parrafos))):
                if re.fullmatch(r"\s*\d+\.\s*", texto_de(parrafos[j])):
                    items.append(j)
                elif items:
                    break
            anexos = {"titulo": i, "items": items}
            break

    casillas = None
    if "tipo_informe" in parr:
        t = parr["tipo_informe"]["indice"]
        desde = parr.get("fecha_elaboracion", {}).get("indice", max(t - 1, 0))
        if _contar_cajas_vacias(parrafos[desde:t + 1]) >= 2:
            casillas = {"desde": desde, "hasta": t, "parcial": 0, "final": 1}

    return {"version": 1, "campos": campos, "parrafos": parr, "casillas": casillas,
            "ejecucion": ejecucion, "anexos": anexos}


def faltantes(mapa: dict) -> list[str]:
    salida = [c for c in OBLIGATORIOS if c != "ejecucion" and c not in (mapa.get("campos") or {})]
    if not mapa.get("ejecucion"):
        salida.append("ejecucion")
    if "periodo" not in (mapa.get("parrafos") or {}):
        salida.append("periodo")
    return salida


# ------------------------------------------------------------------ apoyo de IA
class _Celda(BaseModel):
    tabla: int
    fila: int
    col: int


class _Ejecucion(BaseModel):
    tabla: int
    fila_encabezado: int
    col_item: Optional[int] = None
    col_obligacion: int
    col_actividades: int


class _MapaIA(BaseModel):
    campos: dict[str, Optional[_Celda]] = {}
    ejecucion: Optional[_Ejecucion] = None
    parrafo_fecha_elaboracion: Optional[int] = None
    parrafo_tipo_informe: Optional[int] = None
    parrafo_periodo: Optional[int] = None


def estructura_texto(contenido: bytes, max_lineas: int = 220) -> str:
    doc = Document(io.BytesIO(contenido))
    body = doc.element.body
    lineas = []
    for ti, tbl in enumerate(tablas_cuerpo(body)):
        lineas.append(f"TABLA {ti}:")
        for ri, tr in enumerate(filas(tbl)):
            partes = [f'[{ci}] "{texto_de(tc).strip()[:70]}"' for ci, tc in enumerate(celdas(tr))]
            lineas.append(f"  fila {ri}: " + " | ".join(partes))
    for pi, p in enumerate(parrafos_cuerpo(body)):
        t = texto_de(p).strip()
        if t:
            lineas.append(f'PÁRRAFO {pi}: "{t[:120]}"')
    return "\n".join(lineas[:max_lineas])


def completar_con_ia(contenido: bytes, mapa: dict, user_id: str | None) -> tuple[dict, str | None]:
    falta = faltantes(mapa)
    if not falta or not ia_configurada():
        return mapa, None
    doc = Document(io.BytesIO(contenido))
    body = doc.element.body
    tablas, parrafos = tablas_cuerpo(body), parrafos_cuerpo(body)
    try:
        r, modelo = llamar_ia("mapear_plantilla", [
            {"role": "system", "content": prompts.SISTEMA_MAPA},
            {"role": "user", "content": prompts.usuario_mapa(estructura_texto(contenido), falta)},
        ], _MapaIA, user_id=user_id, max_tokens=1500, temperatura=0.1)
    except IAError as exc:
        log.info("La IA no pudo completar el mapa: %s", str(exc)[:200])
        return mapa, None

    def celda_valida(c: _Celda) -> bool:
        return (0 <= c.tabla < len(tablas) and 0 <= c.fila < len(filas(tablas[c.tabla]))
                and 0 <= c.col < len(celdas(filas(tablas[c.tabla])[c.fila])))

    for campo, celda in r.campos.items():
        if campo in CAMPOS and campo not in mapa["campos"] and celda and celda_valida(celda):
            tr = filas(tablas[celda.tabla])[celda.fila]
            etiqueta = texto_de(celdas(tr)[celda.col - 1]).strip() if celda.col > 0 else ""
            mapa["campos"][campo] = {**celda.model_dump(), "etiqueta": etiqueta, "origen": "ia"}
    e = r.ejecucion
    if not mapa.get("ejecucion") and e and 0 <= e.tabla < len(tablas) \
            and e.fila_encabezado + 1 < len(filas(tablas[e.tabla])):
        mapa["ejecucion"] = {"tabla": e.tabla, "fila_encabezado": e.fila_encabezado,
                             "fila_inicio": e.fila_encabezado + 1, "col_item": e.col_item,
                             "col_obligacion": e.col_obligacion, "col_actividades": e.col_actividades,
                             "fila_formato": None, "eliminar_sobrantes": True, "origen": "ia"}
    for clave, indice in (("fecha_elaboracion", r.parrafo_fecha_elaboracion),
                          ("tipo_informe", r.parrafo_tipo_informe), ("periodo", r.parrafo_periodo)):
        if clave not in mapa["parrafos"] and indice is not None and 0 <= indice < len(parrafos):
            mapa["parrafos"][clave] = {"indice": indice, "texto": texto_de(parrafos[indice]).strip()[:80],
                                       "origen": "ia"}
    return mapa, modelo


def resumen(mapa: dict) -> list[dict]:
    """Lista legible para que el usuario confirme el mapa."""
    filas_resumen = []
    for campo, (nombre, _) in CAMPOS.items():
        ref = (mapa.get("campos") or {}).get(campo)
        filas_resumen.append({
            "campo": campo, "nombre": nombre, "ok": bool(ref),
            "ubicacion": (f"Tabla {ref['tabla'] + 1}, fila {ref['fila'] + 1}: «{ref.get('etiqueta', '')}»"
                          if ref else "No encontrado (quedará en blanco)"),
        })
    e = mapa.get("ejecucion")
    filas_resumen.append({"campo": "ejecucion", "nombre": "Tabla de obligaciones y actividades", "ok": bool(e),
                          "ubicacion": f"Tabla {e['tabla'] + 1}, desde la fila {e['fila_inicio'] + 1}" if e
                          else "No encontrada"})
    nombres = {"fecha_elaboracion": "Fecha de elaboración", "tipo_informe": "Casillas parcial / final",
               "periodo": "Párrafo del periodo"}
    for clave, nombre in nombres.items():
        ref = (mapa.get("parrafos") or {}).get(clave)
        filas_resumen.append({"campo": clave, "nombre": nombre, "ok": bool(ref),
                              "ubicacion": f"«{ref['texto'][:60]}…»" if ref else "No encontrado"})
    a = mapa.get("anexos")
    filas_resumen.append({"campo": "anexos", "nombre": "Lista de anexos", "ok": bool(a and a.get("items")),
                          "ubicacion": f"{len(a['items'])} renglones numerados" if a and a.get("items")
                          else "No encontrada"})
    return filas_resumen

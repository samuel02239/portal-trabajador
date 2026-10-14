"""El informe generado conserva la plantilla y llena cada campo en su lugar (sección 8)."""
import io
from datetime import date

from docx import Document

from servicios import mapa_plantilla, verificacion
from servicios.docx_util import celdas, filas, parrafos_cuerpo, q, tablas_cuerpo, texto_de
from servicios.generador_docx import DatosInforme, ObligacionInforme, generar, llenar_blancos


def _datos(n_obligaciones=10, tipo="parcial", anexos=None):
    obligs = [ObligacionInforme(i, f"OBLIGACIÓN NÚMERO {i} DE PRUEBA",
                                [f"Actividad A de la obligación {i}.", f"Actividad B de la obligación {i}."]
                                if i in (1, 3) else []) for i in range(1, n_obligaciones + 1)]
    return DatosInforme(
        nombre_contratista="PERSONA DE PRUEBA", cedula="1000000001", numero_contrato="1234",
        fecha_contrato=date(2026, 9, 25), objeto="OBJETO DE PRUEBA", fecha_inicio=date(2026, 10, 1),
        fecha_fin=date(2026, 12, 31), dependencia="OFICINA DE PRUEBA", supervisor_nombre="SUPERVISORA",
        supervisor_cargo="JEFE DE OFICINA", periodo_inicio=date(2026, 10, 1), periodo_fin=date(2026, 10, 31),
        fecha_elaboracion=date(2026, 11, 3), tipo=tipo, obligaciones=obligs, anexos=anexos or [])


def _abrir(contenido):
    body = Document(io.BytesIO(contenido)).element.body
    return tablas_cuerpo(body), parrafos_cuerpo(body)


def test_mapa_detecta_todo(plantilla_bytes):
    mapa = mapa_plantilla.detectar(plantilla_bytes)
    assert mapa_plantilla.faltantes(mapa) == []
    assert len(mapa["campos"]) == 9
    assert mapa["ejecucion"]["fila_inicio"] == 2 and mapa["ejecucion"]["fila_formato"] == 2
    assert mapa["casillas"] and mapa["anexos"]["items"]


def test_informe_completo(plantilla_bytes):
    mapa = mapa_plantilla.detectar(plantilla_bytes)
    docx = generar(plantilla_bytes, mapa, _datos(anexos=["foto.jpg", "acta.pdf", "a", "b", "c", "d"]))
    tablas, parrafos = _abrir(docx)

    info = filas(tablas[0])
    assert texto_de(celdas(info[1])[1]) == "PERSONA DE PRUEBA"
    assert texto_de(celdas(info[5])[1]) == "1000000001"
    assert texto_de(celdas(info[3])[1]) == "25 de septiembre de 2026"
    assert texto_de(celdas(info[9])[1]) == "SUPERVISORA - JEFE DE OFICINA"

    ejec = filas(tablas[1])[2:]
    assert len(ejec) == 10  # 8 filas de la plantilla + 2 clonadas
    assert [texto_de(celdas(f)[0]) for f in ejec] == [str(i) for i in range(1, 11)]
    assert texto_de(celdas(ejec[0])[1]) == "OBLIGACIÓN NÚMERO 1 DE PRUEBA"
    vinetas = [p for p in celdas(ejec[0])[2] if p.tag == q("w:p")]
    assert len(vinetas) == 2 and all(p.find(f"{q('w:pPr')}/{q('w:numPr')}") is not None for p in vinetas)
    assert texto_de(celdas(ejec[1])[2]).strip() == ""  # obligación sin actividades queda vacía
    assert "Nota:" not in texto_de(tablas[1])  # se quitó el texto de instrucciones

    todo = "\n".join(texto_de(p) for p in parrafos)
    assert "03 de noviembre de 2026" in todo
    assert "entre el día 1 del mes de octubre, al día 31 del mes de octubre del año 2026" in todo
    assert "NOMBRE DEL SUPERVISOR" in todo and "NOMBRE DEL CONTRATISTA" in todo  # firmas intactas
    assert "6. d" in todo  # se clonaron renglones de anexos

    assert verificacion.comparar_docx(plantilla_bytes, docx) == []


def _x_en_casillas(contenido) -> list[str]:
    body = Document(io.BytesIO(contenido)).element.body
    cajas = []
    for ac in body.iter(q("mc:AlternateContent")):
        caja = ac.find(f".//{q('wps:txbx')}")
        if caja is None:
            continue
        pos = int(ac.find(f".//{q('wp:positionH')}/{q('wp:posOffset')}").text)
        cajas.append((pos, texto_de(caja.find(q("w:txbxContent"))).strip()))
    return [t for _, t in sorted(cajas)]


def test_casilla_parcial_y_final(plantilla_bytes):
    mapa = mapa_plantilla.detectar(plantilla_bytes)
    assert _x_en_casillas(generar(plantilla_bytes, mapa, _datos(tipo="parcial"))) == ["X", ""]
    assert _x_en_casillas(generar(plantilla_bytes, mapa, _datos(tipo="final"))) == ["", "X"]


def test_sobran_filas_se_eliminan(plantilla_bytes):
    mapa = mapa_plantilla.detectar(plantilla_bytes)
    tablas, _ = _abrir(generar(plantilla_bytes, mapa, _datos(n_obligaciones=3)))
    assert len(filas(tablas[1])) == 2 + 3


def _foto(ancho, alto) -> bytes:
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (ancho, alto), (90, 140, 180)).save(b, "JPEG")
    return b.getvalue()


def test_imagenes_al_final_despues_de_validacion(plantilla_bytes):
    from servicios.generador_docx import ImagenAnexo
    mapa = mapa_plantilla.detectar(plantilla_bytes)
    datos = _datos(anexos=["Registro fotográfico del mantenimiento (obligación 1).", "Acta (obligación 3)."])
    datos.imagenes = [ImagenAnexo(1, "Registro fotográfico del mantenimiento", _foto(4000, 3000)),
                      ImagenAnexo(2, "No es una imagen válida", b"\x89PNG\r\n\x1a\nbasura")]
    docx = generar(plantilla_bytes, mapa, datos)
    body = Document(io.BytesIO(docx)).element.body
    hijos = [e for e in body.iterchildren() if e.tag != q("w:sectPr")]
    ultima_tabla = max(i for i, e in enumerate(hijos) if e.tag == q("w:tbl"))
    assert "VALIDACI" in texto_de(hijos[ultima_tabla - 2]) or "VALIDACI" in texto_de(hijos[ultima_tabla - 1])
    despues = hijos[ultima_tabla + 1:]
    texto_final = "\n".join(texto_de(e) for e in despues)
    assert "REGISTRO FOTOGRÁFICO DE LOS ANEXOS" in texto_final
    assert "Anexo 1. Registro fotográfico del mantenimiento." in texto_final
    assert "Anexo 2" not in texto_final  # la imagen dañada no se pega (queda solo en la lista de anexos)
    dibujos = [d for e in despues for d in e.iter(q("w:drawing"))]
    assert len(dibujos) == 1
    assert verificacion.comparar_docx(plantilla_bytes, docx) == []


def test_texto_con_anexos():
    from routers.informes import texto_con_anexos
    assert texto_con_anexos("Realicé el soporte.", []) == "Realicé el soporte."
    assert texto_con_anexos("Realicé el soporte", [(1, "Registro fotográfico del soporte")]) == \
        "Realicé el soporte. Se adjunta como evidencia el Anexo 1: registro fotográfico del soporte."
    assert "los anexos (Anexo 2: registro…; Anexo 3: documento…)" in texto_con_anexos(
        "Asistí.", [(2, "Registro…"), (3, "Documento…")])


def test_descripcion_de_respaldo_sin_ia():
    from servicios.evidencias import descripcion_respaldo, tipo_de
    assert tipo_de("foto.JPG") == "imagen" and tipo_de("acta.pdf") == "documento" and tipo_de("x.xlsx") == "hoja de cálculo"
    d = descripcion_respaldo("Realicé el mantenimiento de los equipos, en desarrollo del soporte técnico.", "imagen")
    assert d == "Registro fotográfico de la actividad: realicé el mantenimiento de los equipos"


def test_llenar_blancos_repartidos_en_runs():
    from lxml import etree
    from servicios.docx_util import NS
    p = etree.fromstring(
        f'<w:p xmlns:w="{NS["w"]}"><w:r><w:t>día ___</w:t></w:r><w:r><w:t>____ del mes </w:t></w:r>'
        '<w:r><w:t>de ____, año __</w:t></w:r></w:p>')
    assert llenar_blancos(p, ["5", "octubre", "2026"]) == 3
    assert texto_de(p) == "día 5 del mes de octubre, año 2026"
